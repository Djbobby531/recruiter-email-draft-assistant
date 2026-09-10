"""
Resume text extraction, used once at upload/replace time (RULE H - never
re-parsed on every incoming email). Supports PDF, modern Word (.docx), and
legacy Word (.doc).
"""
from __future__ import annotations

import re
from pathlib import Path

import docx
from pypdf import PdfReader

# Legacy binary .doc has no reliable pure-Python parser (python-docx only
# reads the modern OOXML .docx format), and this project deliberately avoids
# adding a system-binary dependency (antiword/catdoc/textract) just to read
# an old file format. Instead, runs of printable ASCII embedded in the binary
# are pulled out as a best-effort extraction - imperfect, but it's enough to
# feed the resume matcher's keyword scoring for most simple resumes, and any
# resume where it comes back too sparse is safely marked INDEXED_EMPTY_TEXT
# (routing to manual review) rather than silently pretending to have parsed it.
_PRINTABLE_RUN_RE = re.compile(rb"[ -~]{4,}")


def extract_pdf_text(path: str | Path) -> str:
    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n".join(pages).strip()


def extract_docx_text(path: str | Path) -> str:
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts).strip()


def extract_doc_text(path: str | Path) -> str:
    raw = Path(path).read_bytes()
    chunks = [m.group().decode("ascii", errors="ignore") for m in _PRINTABLE_RUN_RE.finditer(raw)]
    return "\n".join(chunks).strip()


def extract_resume_text(path: str | Path, filename: str) -> str:
    """Dispatches on file extension (not the raw content sniffing) - the
    caller already validated the extension at upload time."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_text(path)
    if suffix == ".docx":
        return extract_docx_text(path)
    if suffix == ".doc":
        return extract_doc_text(path)
    raise ValueError(f"Unsupported resume file type: {suffix}")
