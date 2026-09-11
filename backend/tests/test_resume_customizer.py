"""
Resume match scoring + truthful, format-preserving customization.

`evaluate_and_customize` (the match score/explanation carried onto the
Application) is purely deterministic and ignores any `ai_provider` it's
given. `generate_llm_customized_resume` is the enhanced, opt-in LLM-assisted
CONTENT customization path (header role, summary, skills, experience
bullets) - it never picks WHICH resume is used (that's still the
deterministic resume_matcher), and every field of its output is
independently re-validated against the candidate's own verified resume
metadata before anything is trusted, falling back to `None` (and the caller
falling back further to the plain deterministic `customize_resume_file`) on
any doubt. The core guarantee under test everywhere here: the customizer can
only ever surface skills/facts the candidate's OWN resume metadata already
claims - it must never introduce a skill/technology/experience that isn't
already there, even when an AI provider suggests one.
"""
from __future__ import annotations

import docx

from app.config import Settings
from app.services.resume_customizer import (
    customize_resume_file,
    evaluate_and_customize,
    find_underemphasized_overlap_skills,
    generate_llm_customized_resume,
)
from app.services.resume_matcher import ResumeScore


class _StubResume:
    def __init__(self, filename, file_path, extracted_text, skills):
        self.id = 1
        self.filename = filename
        self.file_path = file_path
        self.extracted_text = extracted_text
        self.extracted_metadata = {"skills": skills}


class _StubAI:
    """Instrumented only to prove it is NEVER called - resume customization
    must stay purely deterministic even when an ai_provider is passed in."""

    def __init__(self, result=None, raise_exc=None):
        self.result = result or {}
        self.raise_exc = raise_exc
        self.called = False

    def evaluate_and_customize_resume(self, resume_text, jd_title, jd_text, candidate_existing_skills):
        self.called = True
        if self.raise_exc:
            raise self.raise_exc
        return self.result


BASE_SCORE = ResumeScore(resume_id=1, filename="r.docx", score=42.0, explanation="baseline deterministic score")


# --- find_underemphasized_overlap_skills -----------------------------------


def test_finds_skills_candidate_has_but_didnt_write_into_prose():
    resume = _StubResume(
        "r.docx", "/tmp/r.docx", extracted_text="I am a data engineer with Python experience.",
        skills={"data_engineering_tools": ["airflow"], "languages": ["python"]},
    )
    result = find_underemphasized_overlap_skills(resume, jd_requirements=["airflow", "python", "kafka"])
    assert result == ["airflow"]  # python already in prose; kafka isn't a candidate skill at all


def test_never_returns_a_skill_the_candidate_does_not_actually_have():
    resume = _StubResume("r.docx", "/tmp/r.docx", extracted_text="", skills={"languages": ["python"]})
    result = find_underemphasized_overlap_skills(resume, jd_requirements=["kubernetes", "terraform"])
    assert result == []  # JD wants these, but candidate's own metadata never claims them


def test_returns_empty_when_everything_already_in_prose():
    resume = _StubResume(
        "r.docx", "/tmp/r.docx", extracted_text="Skilled in Databricks and SQL.",
        skills={"data_engineering_tools": ["databricks"], "languages": ["sql"]},
    )
    result = find_underemphasized_overlap_skills(resume, jd_requirements=["databricks", "sql"])
    assert result == []


def test_skill_buried_outside_summary_is_still_underemphasized():
    """Regression test for a real bug: metadata.skills is itself extracted
    by scanning the WHOLE resume text, so a skill in metadata is always
    present *somewhere* in the full text by construction - checking the
    whole text again (the old behavior) could never find anything
    "underemphasized" and this feature silently never fired for any real
    resume. It must instead check prominence within the Summary section
    specifically: a skill mentioned only in an old job's bullet point, but
    never in the Summary, is still a genuine gap worth surfacing."""
    resume = _StubResume(
        "r.docx", "/tmp/r.docx",
        extracted_text=(
            "Summary\n"
            "10+ years building cloud-native data pipelines with Python and SQL.\n"
            "Experience\n"
            "Data Engineer, Acme Corp\n"
            "Used Databricks to build a one-off migration pipeline in 2019.\n"
        ),
        skills={"data_engineering_tools": ["databricks"], "languages": ["python", "sql"]},
    )
    result = find_underemphasized_overlap_skills(resume, jd_requirements=["databricks", "python"])
    # databricks is genuinely true of the candidate (in metadata) and JD-relevant,
    # but never mentioned in the Summary - still underemphasized despite
    # appearing in an Experience bullet
    assert result == ["databricks"]
    # python IS in the Summary - correctly excluded
    assert "python" not in result


def test_skill_already_prominent_in_summary_is_not_returned():
    resume = _StubResume(
        "r.docx", "/tmp/r.docx",
        extracted_text=(
            "Summary\n"
            "Data engineer skilled in Databricks, Python, and SQL.\n"
            "Experience\n"
            "Built pipelines.\n"
        ),
        skills={"data_engineering_tools": ["databricks"], "languages": ["python"]},
    )
    result = find_underemphasized_overlap_skills(resume, jd_requirements=["databricks", "python"])
    assert result == []


def test_falls_back_to_whole_text_when_no_summary_section_is_identifiable():
    """A resume with no recognizable Summary heading at all keeps the old,
    whole-text-based behavior rather than treating everything as missing."""
    resume = _StubResume(
        "r.docx", "/tmp/r.docx", extracted_text="Just a plain resume mentioning Databricks somewhere.",
        skills={"data_engineering_tools": ["databricks"]},
    )
    result = find_underemphasized_overlap_skills(resume, jd_requirements=["databricks"])
    assert result == []  # already present in the (only) text available to check


# --- evaluate_and_customize (deterministic path, no AI) --------------------


def test_deterministic_path_used_when_no_ai_provider():
    resume = _StubResume(
        "r.docx", "/tmp/r.docx", extracted_text="Data engineer.",
        skills={"data_engineering_tools": ["airflow"]},
    )
    result = evaluate_and_customize(
        resume, jd_title="Data Engineer", jd_text="Requires Airflow", jd_requirements=["airflow"],
        base_score=BASE_SCORE, ai_provider=None,
    )
    assert result.source == "deterministic"
    assert result.match_score == BASE_SCORE.score
    assert "Airflow" in result.additional_points[0]


def test_deterministic_path_has_no_additional_points_when_nothing_missing():
    resume = _StubResume("r.docx", "/tmp/r.docx", extracted_text="Python developer.", skills={"languages": ["python"]})
    result = evaluate_and_customize(
        resume, jd_title="Dev", jd_text="Python role", jd_requirements=["python"],
        base_score=BASE_SCORE, ai_provider=None,
    )
    assert result.additional_points == []


# --- evaluate_and_customize never calls the AI provider --------------------


def test_ai_provider_is_never_consulted_even_when_configured():
    """By explicit configuration choice: resume customization is purely
    deterministic. Even when a (mis)configured ai_provider is passed in, it
    must never be called, and its output must never influence the result -
    not the score, not the explanation, not the additional points."""
    resume = _StubResume(
        "r.docx", "/tmp/r.docx", extracted_text="Data engineer.",
        skills={"data_engineering_tools": ["airflow"]},
    )
    ai = _StubAI({
        "match_score": 91.0, "match_explanation": "Strong airflow alignment",
        "additional_points": ["Hands-on experience orchestrating pipelines with Airflow"],
    })
    result = evaluate_and_customize(
        resume, jd_title="Data Engineer", jd_text="Requires Airflow", jd_requirements=["airflow"],
        base_score=BASE_SCORE, ai_provider=ai,
    )
    assert ai.called is False
    assert result.source == "deterministic"
    assert result.match_score == BASE_SCORE.score
    assert result.match_explanation == BASE_SCORE.explanation
    assert result.additional_points == ["Additional Relevant Skills: Airflow"]


def test_ai_provider_raising_is_irrelevant_since_it_is_never_called():
    resume = _StubResume(
        "r.docx", "/tmp/r.docx", extracted_text="Data engineer.",
        skills={"data_engineering_tools": ["airflow"]},
    )
    ai = _StubAI(raise_exc=ConnectionError("ollama down"))
    result = evaluate_and_customize(
        resume, jd_title="Data Engineer", jd_text="Requires Airflow", jd_requirements=["airflow"],
        base_score=BASE_SCORE, ai_provider=ai,
    )
    assert ai.called is False
    assert result.source == "deterministic"
    assert result.match_score == BASE_SCORE.score


# --- customize_resume_file (format-preserving DOCX only) -------------------


def _make_docx_resume(tmp_path, filename="resume.docx", paragraphs=None):
    document = docx.Document()
    for text in paragraphs or ["Diwakar Jilakara", "Senior Data Engineer"]:
        document.add_paragraph(text)
    path = tmp_path / filename
    document.save(str(path))
    return _StubResume(filename, str(path), extracted_text="\n".join(paragraphs or []), skills={})


def test_customize_resume_file_returns_none_when_no_additional_points(tmp_path):
    resume = _make_docx_resume(tmp_path)
    assert customize_resume_file(resume, [], dest_dir=str(tmp_path)) is None


def test_customize_resume_file_returns_none_for_pdf(tmp_path):
    resume = _StubResume("resume.pdf", str(tmp_path / "resume.pdf"), extracted_text="", skills={})
    assert customize_resume_file(resume, ["Additional Relevant Skills: Airflow"], dest_dir=str(tmp_path)) is None


def test_customize_resume_file_returns_none_for_doc(tmp_path):
    resume = _StubResume("resume.doc", str(tmp_path / "resume.doc"), extracted_text="", skills={})
    assert customize_resume_file(resume, ["Additional Relevant Skills: Airflow"], dest_dir=str(tmp_path)) is None


def test_customize_resume_file_creates_new_docx_without_touching_original(tmp_path):
    original_dir = tmp_path / "originals"
    original_dir.mkdir()
    dest_dir = tmp_path / "customized"
    resume = _make_docx_resume(original_dir, paragraphs=["Diwakar Jilakara", "Senior Data Engineer"])
    original_text_before = docx.Document(resume.file_path).paragraphs[0].text

    result = customize_resume_file(resume, ["Additional Relevant Skills: Airflow, Kafka"], dest_dir=str(dest_dir))

    assert result is not None
    new_path, display_filename = result
    assert new_path != resume.file_path
    assert display_filename == "Customized_resume.docx"

    # original untouched
    assert docx.Document(resume.file_path).paragraphs[0].text == original_text_before
    original_paragraph_count = len(docx.Document(resume.file_path).paragraphs)

    # new file has the original content PLUS the new section
    new_document = docx.Document(new_path)
    new_texts = [p.text for p in new_document.paragraphs]
    assert "Diwakar Jilakara" in new_texts
    assert "Additional Relevant Skills" in new_texts
    assert any("Airflow" in t for t in new_texts)
    assert len(new_document.paragraphs) > original_paragraph_count


# --- customize_resume_file: section-aware insertion (Experience/Summary) ---


def _make_structured_docx_resume(tmp_path, filename="resume.docx"):
    """A resume shaped like a real one: header, Summary (prose), Experience
    (bulleted), Education - so section-boundary detection has something
    realistic to work with."""
    document = docx.Document()
    document.add_paragraph("Diwakar Jilakara")
    document.add_paragraph("Senior Data Engineer")
    document.add_paragraph("Summary")
    document.add_paragraph("Experienced data engineer with a background in cloud data platforms.")
    document.add_paragraph("Experience")
    document.add_paragraph("Built and maintained ETL pipelines using Airflow.", style="List Bullet")
    document.add_paragraph("Optimized Spark jobs for large-scale processing.", style="List Bullet")
    document.add_paragraph("Education")
    document.add_paragraph("B.S. Computer Science")
    path = tmp_path / filename
    document.save(str(path))
    return _StubResume(filename, str(path), extracted_text="", skills={})


def test_customize_resume_file_inserts_points_into_experience_before_education(tmp_path):
    resume = _make_structured_docx_resume(tmp_path)
    result = customize_resume_file(
        resume, ["Implemented CI/CD pipelines using Jenkins and Terraform"], dest_dir=str(tmp_path / "out"),
    )
    assert result is not None
    new_path, _ = result
    new_texts = [p.text for p in docx.Document(new_path).paragraphs]

    # inserted, and landed BEFORE "Education" (i.e. still inside Experience),
    # not appended at the very end of the document
    assert any("Jenkins" in t for t in new_texts)
    experience_idx = new_texts.index("Experience")
    education_idx = new_texts.index("Education")
    jenkins_idx = next(i for i, t in enumerate(new_texts) if "Jenkins" in t)
    assert experience_idx < jenkins_idx < education_idx
    # no separate "Additional Relevant Skills" appendix was needed
    assert "Additional Relevant Skills" not in new_texts


def test_customize_resume_file_preserves_bullet_style_in_experience(tmp_path):
    resume = _make_structured_docx_resume(tmp_path)
    result = customize_resume_file(
        resume, ["Automated deployments with Terraform"], dest_dir=str(tmp_path / "out"),
    )
    new_path, _ = result
    new_document = docx.Document(new_path)
    inserted = next(p for p in new_document.paragraphs if "Terraform" in p.text)
    existing_bullet = next(p for p in new_document.paragraphs if "Airflow" in p.text)
    assert inserted.style.name == existing_bullet.style.name == "List Bullet"


def test_customize_resume_file_does_not_modify_any_existing_paragraph_text(tmp_path):
    resume = _make_structured_docx_resume(tmp_path)
    before_texts = [p.text for p in docx.Document(resume.file_path).paragraphs]

    result = customize_resume_file(resume, ["A new point"], dest_dir=str(tmp_path / "out"))
    new_path, _ = result
    new_texts = [p.text for p in docx.Document(new_path).paragraphs]

    # every original paragraph's text still appears, unmodified, in order
    filtered = [t for t in new_texts if t in before_texts]
    assert filtered == before_texts
    # and the original file on disk is completely untouched
    assert [p.text for p in docx.Document(resume.file_path).paragraphs] == before_texts


def test_customize_resume_file_falls_back_to_summary_when_no_experience_heading(tmp_path):
    document = docx.Document()
    document.add_paragraph("Diwakar Jilakara")
    document.add_paragraph("Summary")
    document.add_paragraph("Experienced engineer.")
    document.add_paragraph("Education")
    document.add_paragraph("B.S. Computer Science")
    path = tmp_path / "resume.docx"
    document.save(str(path))
    resume = _StubResume("resume.docx", str(path), extracted_text="", skills={})

    result = customize_resume_file(resume, ["Skilled in Kubernetes-based deployments"], dest_dir=str(tmp_path / "out"))
    new_texts = [p.text for p in docx.Document(result[0]).paragraphs]

    assert any("Kubernetes" in t for t in new_texts)
    summary_idx = new_texts.index("Summary")
    education_idx = new_texts.index("Education")
    kube_idx = next(i for i, t in enumerate(new_texts) if "Kubernetes" in t)
    assert summary_idx < kube_idx < education_idx
    assert "Additional Relevant Skills" not in new_texts


def test_customize_resume_file_falls_back_to_appendix_when_no_recognized_section(tmp_path):
    """A resume with no detectable Summary/Experience heading at all must
    still get the points added - just via the old safe appendix, not
    guessed insertion."""
    resume = _make_docx_resume(tmp_path, paragraphs=["Diwakar Jilakara", "Senior Data Engineer"])
    result = customize_resume_file(resume, ["Skilled in Kubernetes"], dest_dir=str(tmp_path / "out"))
    new_texts = [p.text for p in docx.Document(result[0]).paragraphs]
    assert "Additional Relevant Skills" in new_texts
    assert any("Kubernetes" in t for t in new_texts)


# --- post-save formatting-integrity check: every failure reverts to the ---
# --- original, unmodified resume, for every email                      ---


def test_is_subsequence_helper_accepts_original_plus_insertions():
    from app.services.resume_customizer import _is_subsequence_of_texts
    original = ["A", "B", "C"]
    new = ["A", "INSERTED", "B", "C", "INSERTED 2"]
    assert _is_subsequence_of_texts(original, new) is True


def test_is_subsequence_helper_rejects_modified_text():
    from app.services.resume_customizer import _is_subsequence_of_texts
    original = ["A", "B", "C"]
    new = ["A", "B (edited)", "C"]  # "B" itself was changed, not just added-around
    assert _is_subsequence_of_texts(original, new) is False


def test_is_subsequence_helper_rejects_reordered_text():
    from app.services.resume_customizer import _is_subsequence_of_texts
    original = ["A", "B", "C"]
    new = ["A", "C", "B"]  # same texts, wrong order
    assert _is_subsequence_of_texts(original, new) is False


def test_is_subsequence_helper_rejects_dropped_paragraph():
    from app.services.resume_customizer import _is_subsequence_of_texts
    original = ["A", "B", "C"]
    new = ["A", "C"]  # "B" went missing
    assert _is_subsequence_of_texts(original, new) is False


def test_customization_is_format_safe_catches_a_content_losing_save(tmp_path):
    """Direct unit test of the integrity check: a saved file with LESS
    content than the original (simulating a corrupted/lossy save) must be
    rejected."""
    from app.services.resume_customizer import _customization_is_format_safe

    original = docx.Document()
    original.add_paragraph("Diwakar Jilakara")
    original.add_paragraph("Senior Data Engineer")
    original_path = tmp_path / "original.docx"
    original.save(str(original_path))

    broken = docx.Document()
    broken.add_paragraph("Only this survived")  # original content lost
    broken_path = tmp_path / "broken.docx"
    broken.save(str(broken_path))

    assert _customization_is_format_safe(docx.Document(str(original_path)), str(broken_path)) is False


def test_customize_resume_file_reverts_to_original_and_cleans_up_when_integrity_check_fails(tmp_path, monkeypatch):
    """Integration behavior: when the post-save integrity check fails, the
    customization must be discarded entirely (return None -> caller attaches
    the plain original resume) and the bad file must not be left behind."""
    import app.services.resume_customizer as resume_customizer_module

    resume = _make_structured_docx_resume(tmp_path)
    out_dir = tmp_path / "out"

    monkeypatch.setattr(resume_customizer_module, "_customization_is_format_safe", lambda *a, **k: False)

    result = customize_resume_file(resume, ["Skilled in Kubernetes"], dest_dir=str(out_dir))
    assert result is None
    assert list(out_dir.glob("*.docx")) == []  # the rejected file was cleaned up, not left behind


def test_customize_resume_file_reverts_to_original_when_document_open_raises(tmp_path, monkeypatch):
    """A resume file that can't even be opened must fail safe (None), not
    raise out of the pipeline."""
    import app.services.resume_customizer as resume_customizer_module

    resume = _make_structured_docx_resume(tmp_path)

    def _raise(*args, **kwargs):
        raise ValueError("simulated corrupt docx")

    monkeypatch.setattr(resume_customizer_module.docx, "Document", _raise)

    result = customize_resume_file(resume, ["Skilled in Kubernetes"], dest_dir=str(tmp_path / "out"))
    assert result is None


def test_customize_resume_file_keeps_a_genuinely_clean_edit(tmp_path):
    """Sanity check: the integrity check doesn't reject normal, correct
    edits - only actually-broken ones."""
    resume = _make_structured_docx_resume(tmp_path)
    result = customize_resume_file(resume, ["Skilled in Kubernetes-based deployments"], dest_dir=str(tmp_path / "out"))
    assert result is not None


# =============================================================================
# generate_llm_customized_resume - header role, summary, skills placed into
# an existing category, experience bullets. Ollama-shaped stub throughout;
# every unsafe/unverified LLM output must be rejected and the file's
# formatting/structure must survive.
# =============================================================================


class _PlanAI:
    def __init__(self, plan=None, raises=None):
        self.plan = plan if plan is not None else {}
        self.raises = raises
        self.calls = []

    def generate_resume_customization_plan(
        self, resume_text, resume_structure, jd_title, jd_text, jd_requirements,
        approved_skills, approved_experience_identifiers,
    ):
        self.calls.append({
            "resume_structure": resume_structure, "approved_skills": approved_skills,
            "approved_experience_identifiers": approved_experience_identifiers,
        })
        if self.raises:
            raise self.raises
        return self.plan


def _llm_settings(**overrides):
    kwargs = dict(
        RESUME_LLM_CUSTOMIZATION_ENABLED=True, RESUME_LLM_MAX_SUMMARY_POINTS=2,
        RESUME_LLM_MAX_EXPERIENCE_POINTS=2, RESUME_LLM_MAX_SKILLS=6,
    )
    kwargs.update(overrides)
    return Settings(**kwargs)


def _make_rich_docx_resume(tmp_path, filename="resume.docx", header_role="Senior Data Engineer"):
    """Header role + Skills (paragraph-based categories) + Experience (two
    companies, each with bullets) + Education - the shape the acceptance
    example in the spec describes."""
    document = docx.Document()
    document.add_paragraph("Diwakar Jilakara")
    role_p = document.add_paragraph(header_role)
    role_p.runs[0].bold = True
    document.add_paragraph("Summary")
    document.add_paragraph("Experienced data engineer with a background in cloud data platforms.")
    document.add_paragraph("Skills")
    document.add_paragraph("Cloud: AWS, Azure")
    document.add_paragraph("Programming Languages: Python, Java")
    document.add_paragraph("Big Data: Spark, Kafka")
    document.add_paragraph("Databases: Snowflake, Oracle")
    document.add_paragraph("Experience")
    document.add_paragraph("Company A")
    document.add_paragraph("Senior Data Engineer")
    document.add_paragraph("Built and maintained ETL pipelines.", style="List Bullet")
    document.add_paragraph("Optimized Spark jobs for large-scale processing.", style="List Bullet")
    document.add_paragraph("Company B")
    document.add_paragraph("Data Engineer")
    document.add_paragraph("Maintained legacy reporting pipelines.", style="List Bullet")
    document.add_paragraph("Education")
    document.add_paragraph("B.S. Computer Science")
    path = tmp_path / filename
    document.save(str(path))
    return _StubResume(
        filename, str(path),
        extracted_text="Cloud: AWS, Azure\nBuilt Airflow-based orchestration at Company A.",
        skills={
            "cloud_platforms": ["aws", "azure"], "languages": ["python", "java"],
            "data_engineering_tools": ["spark", "kafka", "airflow"],
            "databases": ["snowflake", "oracle"],
        },
    )


def _rich_resume_titles(resume, header_role="Senior Data Engineer"):
    resume.extracted_metadata["job_titles"] = [header_role]
    return resume


def test_llm_customization_none_when_ai_provider_is_none(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    result = generate_llm_customized_resume(
        resume, jd_title="Senior AWS Data Engineer", jd_text="AWS", jd_requirements=["aws"],
        ai_provider=None, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is None


def test_llm_customization_none_when_disabled_by_config(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={"header_role": "Senior AWS Data Engineer"})
    result = generate_llm_customized_resume(
        resume, jd_title="Senior AWS Data Engineer", jd_text="AWS", jd_requirements=["aws"],
        ai_provider=ai, settings=_llm_settings(RESUME_LLM_CUSTOMIZATION_ENABLED=False),
        dest_dir=str(tmp_path / "out"),
    )
    assert result is None
    assert ai.calls == []


def test_llm_customization_none_for_pdf(tmp_path):
    resume = _StubResume("resume.pdf", str(tmp_path / "resume.pdf"), extracted_text="", skills={})
    ai = _PlanAI(plan={"header_role": "Senior AWS Data Engineer"})
    result = generate_llm_customized_resume(
        resume, jd_title="Senior AWS Data Engineer", jd_text="AWS", jd_requirements=["aws"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is None
    assert ai.calls == []


def test_llm_customization_header_role_replaced_preserving_run_formatting(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={"header_role": "Senior AWS Data Engineer", "summary_points": [], "skills_to_add": [], "experience_updates": []})
    result = generate_llm_customized_resume(
        resume, jd_title="Senior AWS Data Engineer", jd_text="AWS role", jd_requirements=["aws"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is not None
    new_path, filename = result
    assert filename == "diwakar_Senior_AWS_Data_Engineer.docx"

    new_document = docx.Document(new_path)
    role_paragraph = next(p for p in new_document.paragraphs if p.text == "Senior AWS Data Engineer")
    assert role_paragraph.runs[0].bold is True  # original bold formatting preserved
    assert "Diwakar Jilakara" in [p.text for p in new_document.paragraphs]  # name untouched

    # original library file is completely untouched
    original_texts = [p.text for p in docx.Document(resume.file_path).paragraphs]
    assert "Senior Data Engineer" in original_texts
    assert "Senior AWS Data Engineer" not in original_texts


def test_llm_customization_header_role_strips_hash_and_symbols(tmp_path):
    """A JD-derived role can carry decorative junk (#, emoji) forward from a
    messy recruiter subject line - even though jd_extractor already cleans
    job_title at the source, the AI's own header_role suggestion must be
    cleaned independently too (defense in depth) before it ever lands in the
    resume header or the output filename."""
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={
        "header_role": "🚨 Senior #AWS Data Engineer", "summary_points": [], "skills_to_add": [], "experience_updates": [],
    })
    result = generate_llm_customized_resume(
        resume, jd_title="Senior AWS Data Engineer", jd_text="AWS role", jd_requirements=["aws"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is not None
    new_path, filename = result
    new_texts = [p.text for p in docx.Document(new_path).paragraphs]
    assert "Senior AWS Data Engineer" in new_texts
    assert not any("#" in t or "🚨" in t for t in new_texts if t != "Senior Data Engineer")
    assert "#" not in filename
    assert filename == "diwakar_Senior_AWS_Data_Engineer.docx"


def test_sanitize_filename_component_normalizes_typographic_dashes_to_hyphen():
    from app.services.resume_customizer import _sanitize_filename_component

    assert _sanitize_filename_component("Senior Data Engineer – dbt") == "Senior_Data_Engineer_-_dbt"
    assert _sanitize_filename_component("Data Engineer — AWS") == "Data_Engineer_-_AWS"


def test_llm_customization_header_role_found_via_heuristic_when_no_exact_title_match(tmp_path):
    """Real-world regression: a table-based header with Name / role-tagline /
    contact-info stacked as separate cell paragraphs, where the tagline
    ("Azure Data Engineer - Databricks, ADF & Data Security") doesn't
    exactly match any of the resume's auto-detected job_titles (which found
    a plainer "Senior Azure Data Engineer" from an Experience entry
    elsewhere). The heuristic fallback must still find and replace it."""
    document = docx.Document()
    table = document.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.paragraphs[0].text = "Diwakar J"
    role_p = cell.add_paragraph("Azure Data Engineer - Databricks, ADF & Data Security")
    role_p.runs[0].bold = True
    cell.add_paragraph("test.candidate@example.com")
    cell.add_paragraph("(555) 123-4567")
    document.add_paragraph("Professional Summary")
    document.add_paragraph("Experienced data engineer.")
    document.add_paragraph("Skills")
    document.add_paragraph("Cloud: AWS, Azure")
    document.add_paragraph("Experience")
    document.add_paragraph("Company A")
    document.add_paragraph("Senior Azure Data Engineer")
    document.add_paragraph("Built pipelines.", style="List Bullet")
    path = tmp_path / "table_header.docx"
    document.save(str(path))
    resume = _StubResume(
        "table_header.docx", str(path), extracted_text="",
        skills={"cloud_platforms": ["aws", "azure"]},
    )
    resume.extracted_metadata["job_titles"] = ["Senior Azure Data Engineer"]  # NOT the tagline text

    ai = _PlanAI(plan={
        "header_role": "Senior AWS Data Engineer", "summary_points": [], "skills_to_add": [], "experience_updates": [],
    })
    result = generate_llm_customized_resume(
        resume, jd_title="Senior AWS Data Engineer", jd_text="AWS role", jd_requirements=["aws"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is not None
    new_document = docx.Document(result[0])
    new_cell_texts = [p.text for p in new_document.tables[0].rows[0].cells[0].paragraphs]
    assert "Senior AWS Data Engineer" in new_cell_texts
    assert "Azure Data Engineer - Databricks, ADF & Data Security" not in new_cell_texts
    # neighboring header lines (name, email, phone) are completely untouched
    assert "Diwakar J" in new_cell_texts
    assert "test.candidate@example.com" in new_cell_texts
    assert "(555) 123-4567" in new_cell_texts
    role_paragraph = next(p for p in new_document.tables[0].rows[0].cells[0].paragraphs if p.text == "Senior AWS Data Engineer")
    assert role_paragraph.runs[0].bold is True


def test_llm_customization_adds_verified_skill_to_correct_category(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={
        "header_role": "", "summary_points": [],
        "skills_to_add": [{"skill": "airflow", "category": "Big Data", "reason": "verified"}],
        "experience_updates": [],
    })
    result = generate_llm_customized_resume(
        resume, jd_title="Data Engineer", jd_text="Airflow required", jd_requirements=["airflow"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is not None
    texts = [p.text for p in docx.Document(result[0]).paragraphs]
    assert "Big Data: Spark, Kafka, Airflow" in texts
    # untouched categories remain byte-for-byte identical
    assert "Cloud: AWS, Azure" in texts
    assert "Programming Languages: Python, Java" in texts


def test_llm_customization_rejects_an_unverified_skill(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={
        "header_role": "", "summary_points": [],
        "skills_to_add": [{"skill": "Terraform", "category": "DevOps", "reason": "fabricated - not verified"}],
        "experience_updates": [],
    })
    result = generate_llm_customized_resume(
        resume, jd_title="Data Engineer", jd_text="Terraform required", jd_requirements=["terraform"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is None  # nothing else in the plan was safe to add either


def test_llm_customization_never_duplicates_an_existing_skill(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={
        "header_role": "", "summary_points": [],
        "skills_to_add": [{"skill": "AWS", "category": "Cloud", "reason": "already listed"}],
        "experience_updates": [],
    })
    result = generate_llm_customized_resume(
        resume, jd_title="Data Engineer", jd_text="AWS required", jd_requirements=["aws"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is None


def test_llm_customization_treats_aws_alias_as_a_duplicate(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={
        "header_role": "", "summary_points": [],
        "skills_to_add": [{"skill": "Amazon Web Services", "category": "Cloud", "reason": "alias of AWS"}],
        "experience_updates": [],
    })
    result = generate_llm_customized_resume(
        resume, jd_title="Data Engineer", jd_text="AWS required", jd_requirements=["aws"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is None


def test_llm_customization_adds_experience_bullets_to_the_matched_company_only(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={
        "header_role": "", "summary_points": [],
        "skills_to_add": [],
        "experience_updates": [{
            "experience_identifier": "Company A",
            "points": ["Automated data pipeline orchestration using Airflow on AWS."],
        }],
    })
    result = generate_llm_customized_resume(
        resume, jd_title="AWS Data Engineer", jd_text="AWS Airflow", jd_requirements=["aws", "airflow"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is not None
    texts = [p.text for p in docx.Document(result[0]).paragraphs]

    company_a_idx = texts.index("Company A")
    company_b_idx = texts.index("Company B")
    bullet_idx = next(i for i, t in enumerate(texts) if "Automated data pipeline orchestration" in t)
    assert company_a_idx < bullet_idx < company_b_idx  # landed inside Company A's block, not Company B's

    # Company A's existing bullets are unchanged, Company B got nothing added
    assert "Built and maintained ETL pipelines." in texts
    assert "Optimized Spark jobs for large-scale processing." in texts
    assert "Maintained legacy reporting pipelines." in texts
    b_bullets_after = texts[company_b_idx:texts.index("Education")]
    assert b_bullets_after == ["Company B", "Data Engineer", "Maintained legacy reporting pipelines."]


def test_llm_customization_caps_experience_bullets_at_configured_maximum(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={
        "header_role": "", "summary_points": [],
        "skills_to_add": [],
        "experience_updates": [{
            "experience_identifier": "Company A",
            "points": [
                "Automated data pipeline orchestration using Airflow.",
                "Optimized Spark-based large-scale data processing workflows.",
                "A third bullet that must never be inserted at all.",
            ],
        }],
    })
    result = generate_llm_customized_resume(
        resume, jd_title="AWS Data Engineer", jd_text="AWS Airflow Spark", jd_requirements=["aws", "airflow", "spark"],
        ai_provider=ai, settings=_llm_settings(RESUME_LLM_MAX_EXPERIENCE_POINTS=2),
        dest_dir=str(tmp_path / "out"),
    )
    assert result is not None
    texts = [p.text for p in docx.Document(result[0]).paragraphs]
    assert not any("must never be inserted" in t for t in texts)
    assert any("orchestration using Airflow" in t for t in texts)
    assert any("Optimized Spark-based" in t for t in texts)


def test_llm_customization_rejects_hedge_language(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={
        "header_role": "", "summary_points": [
            "Exposure to frontend work - not mentioned, but implied through mention of Python.",
        ],
        "skills_to_add": [], "experience_updates": [],
    })
    result = generate_llm_customized_resume(
        resume, jd_title="Data Engineer", jd_text="Python required", jd_requirements=["python"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is None


def test_llm_customization_rejects_an_unapproved_technology_mentioned_in_prose(tmp_path):
    """The sentence itself never says "skills_to_add": ["Kubernetes"] - it just
    slips Kubernetes into free-text prose. Must still be rejected."""
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={
        "header_role": "", "summary_points": [],
        "skills_to_add": [], "experience_updates": [{
            "experience_identifier": "Company A",
            "points": ["Deployed and managed Kubernetes clusters for container orchestration."],
        }],
    })
    result = generate_llm_customized_resume(
        resume, jd_title="Data Engineer", jd_text="Kubernetes required", jd_requirements=["kubernetes"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is None


def test_llm_customization_adds_summary_points_capped_and_matching_format(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={
        "header_role": "", "summary_points": [
            "Senior data engineer experienced with AWS and Python-based pipelines.",
            "Second valid statement.",
            "A third one that must be dropped by the configured cap.",
        ],
        "skills_to_add": [], "experience_updates": [],
    })
    result = generate_llm_customized_resume(
        resume, jd_title="AWS Data Engineer", jd_text="AWS Python", jd_requirements=["aws", "python"],
        ai_provider=ai, settings=_llm_settings(RESUME_LLM_MAX_SUMMARY_POINTS=2), dest_dir=str(tmp_path / "out"),
    )
    assert result is not None
    texts = [p.text for p in docx.Document(result[0]).paragraphs]
    assert "Senior data engineer experienced with AWS and Python-based pipelines." in texts
    assert "Second valid statement." in texts
    assert not any("must be dropped" in t for t in texts)


def test_llm_customization_provider_exception_returns_none(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(raises=ConnectionError("ollama unreachable"))
    result = generate_llm_customized_resume(
        resume, jd_title="AWS Data Engineer", jd_text="AWS", jd_requirements=["aws"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is None


def test_llm_customization_malformed_plan_returns_none(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan="this is not even a dict")
    result = generate_llm_customized_resume(
        resume, jd_title="AWS Data Engineer", jd_text="AWS", jd_requirements=["aws"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is None


def test_llm_customization_empty_plan_returns_none(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={"header_role": "", "summary_points": [], "skills_to_add": [], "experience_updates": []})
    result = generate_llm_customized_resume(
        resume, jd_title="AWS Data Engineer", jd_text="AWS", jd_requirements=["aws"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is None


def test_llm_customization_original_file_untouched_after_customization(tmp_path):
    import os

    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    original_path = resume.file_path
    stat_before = os.stat(original_path)

    ai = _PlanAI(plan={"header_role": "Senior AWS Data Engineer", "summary_points": [], "skills_to_add": [], "experience_updates": []})
    result = generate_llm_customized_resume(
        resume, jd_title="Senior AWS Data Engineer", jd_text="AWS", jd_requirements=["aws"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is not None
    stat_after = os.stat(original_path)
    assert stat_before.st_size == stat_after.st_size
    assert stat_before.st_mtime == stat_after.st_mtime
    assert result[0] != original_path


def test_llm_customization_two_runs_against_different_jds_never_leak_into_each_other(tmp_path):
    """Both runs must start fresh from the ORIGINAL library resume - a
    customization from run 1 must never carry over into run 2's output."""
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))

    ai_aws = _PlanAI(plan={"header_role": "Senior AWS Data Engineer", "summary_points": [], "skills_to_add": [], "experience_updates": []})
    result_aws = generate_llm_customized_resume(
        resume, jd_title="Senior AWS Data Engineer", jd_text="AWS", jd_requirements=["aws"],
        ai_provider=ai_aws, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    ai_azure = _PlanAI(plan={"header_role": "Senior Azure Data Engineer", "summary_points": [], "skills_to_add": [], "experience_updates": []})
    result_azure = generate_llm_customized_resume(
        resume, jd_title="Senior Azure Data Engineer", jd_text="Azure", jd_requirements=["azure"],
        ai_provider=ai_azure, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )

    assert result_aws is not None and result_azure is not None
    aws_texts = [p.text for p in docx.Document(result_aws[0]).paragraphs]
    azure_texts = [p.text for p in docx.Document(result_azure[0]).paragraphs]
    assert "Senior AWS Data Engineer" in aws_texts and "Senior Azure Data Engineer" not in aws_texts
    assert "Senior Azure Data Engineer" in azure_texts and "Senior AWS Data Engineer" not in azure_texts
    assert result_aws[1] == "diwakar_Senior_AWS_Data_Engineer.docx"
    assert result_azure[1] == "diwakar_Senior_Azure_Data_Engineer.docx"


def test_llm_customization_filename_sanitizes_invalid_characters(tmp_path):
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={
        "header_role": 'Data Engineer: AWS/Azure "Lead"', "summary_points": [], "skills_to_add": [], "experience_updates": [],
    })
    result = generate_llm_customized_resume(
        resume, jd_title="Data Engineer", jd_text="AWS", jd_requirements=["aws"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is not None
    _, filename = result
    assert filename.startswith("diwakar_")
    assert filename.endswith(".docx")
    for bad_char in '/\\:*?"<>|':
        assert bad_char not in filename


def test_llm_customization_no_change_when_header_already_matches_jd(tmp_path):
    """If the AI proposes the SAME role the resume already has, nothing
    changes there - only genuinely new content still applies."""
    resume = _rich_resume_titles(_make_rich_docx_resume(tmp_path))
    ai = _PlanAI(plan={
        "header_role": "Senior Data Engineer",  # identical to current
        "summary_points": [], "skills_to_add": [], "experience_updates": [],
    })
    result = generate_llm_customized_resume(
        resume, jd_title="Senior Data Engineer", jd_text="Python", jd_requirements=["python"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is None  # no-op header + nothing else to add -> falls back


def test_llm_customization_no_recognizable_structure_skips_the_llm_call_entirely(tmp_path):
    """A resume with no detectable header/skills/experience/summary at all
    (e.g. empty metadata, no matching title anywhere) shouldn't even bother
    calling the AI provider."""
    document = docx.Document()
    document.add_paragraph("Just some unrelated plain text.")
    path = tmp_path / "bare.docx"
    document.save(str(path))
    resume = _StubResume("bare.docx", str(path), extracted_text="", skills={})

    ai = _PlanAI(plan={"header_role": "Whatever"})
    result = generate_llm_customized_resume(
        resume, jd_title="Data Engineer", jd_text="Python", jd_requirements=["python"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is None
    assert ai.calls == []


def test_llm_customization_skill_added_to_table_cell_category_passes_validation(tmp_path):
    """Real-world regression: a table-based header AND table-based skills
    (2-column rows: label cell | comma-separated items cell) - a very common
    resume layout. A validated skill landing in that TABLE cell paragraph
    used to be applied but never tracked for the post-save integrity check,
    so `_llm_customization_is_valid` always found the cell's OLD text
    "missing" (it now legitimately differs) and discarded the entire,
    otherwise-successful customization - silently falling back to the plain
    deterministic path for every resume that keeps skills in a table. Both
    the header AND the table-cell skill addition must now survive validation
    together."""
    document = docx.Document()
    header_table = document.add_table(rows=1, cols=1)
    header_table.rows[0].cells[0].paragraphs[0].text = "Senior Azure Data Engineer"
    document.add_paragraph("Experience")
    document.add_paragraph("Company A")
    document.add_paragraph("Senior Azure Data Engineer")
    document.add_paragraph("Built pipelines.", style="List Bullet")

    skills_table = document.add_table(rows=1, cols=2)
    row = skills_table.rows[0]
    row.cells[0].paragraphs[0].text = "Languages"
    row.cells[1].paragraphs[0].text = "Python, Java"
    path = tmp_path / "table_skills.docx"
    document.save(str(path))
    resume = _StubResume(
        "table_skills.docx", str(path), extracted_text="",
        skills={"languages": ["python", "java", "go"]},
    )
    resume.extracted_metadata["job_titles"] = ["Senior Azure Data Engineer"]

    ai = _PlanAI(plan={
        "header_role": "Senior GoLang Data Engineer", "summary_points": [],
        "skills_to_add": [{"skill": "go", "category": "Languages", "reason": "verified"}],
        "experience_updates": [],
    })
    result = generate_llm_customized_resume(
        resume, jd_title="Senior GoLang Data Engineer", jd_text="Go required", jd_requirements=["go"],
        ai_provider=ai, settings=_llm_settings(), dest_dir=str(tmp_path / "out"),
    )
    assert result is not None, "a validated table-cell skill addition must not cause the whole plan to be discarded"
    new_document = docx.Document(result[0])
    header_texts = [p.text for p in new_document.tables[0].rows[0].cells[0].paragraphs]
    assert "Senior GoLang Data Engineer" in header_texts
    skills_cell_text = new_document.tables[1].rows[0].cells[1].text
    assert skills_cell_text == "Python, Java, GO"
