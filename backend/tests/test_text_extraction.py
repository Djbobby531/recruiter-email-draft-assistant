"""
Resume format support: PDF (existing), DOCX (python-docx), and legacy DOC
(best-effort printable-text extraction, no external binary dependency).
"""
from __future__ import annotations

import docx
import pytest

from app.utils.text_extraction import extract_doc_text, extract_docx_text, extract_resume_text


def _make_docx(tmp_path, paragraphs, table_rows=None):
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, cell_text in enumerate(row):
                table.cell(r, c).text = cell_text
    path = tmp_path / "resume.docx"
    document.save(str(path))
    return path


def test_extract_docx_text_reads_paragraphs(tmp_path):
    path = _make_docx(tmp_path, ["Diwakar Jilakara", "Senior Data Engineer", "Skills: Python, SQL, Databricks"])
    text = extract_docx_text(path)
    assert "Diwakar Jilakara" in text
    assert "Databricks" in text


def test_extract_docx_text_includes_table_content(tmp_path):
    path = _make_docx(tmp_path, ["Resume"], table_rows=[["Skill", "Years"], ["Python", "8"]])
    text = extract_docx_text(path)
    assert "Python" in text
    assert "8" in text


def test_extract_docx_text_skips_blank_paragraphs(tmp_path):
    path = _make_docx(tmp_path, ["Python", "", "   ", "SQL"])
    text = extract_docx_text(path)
    assert text == "Python\nSQL"


def test_extract_doc_text_pulls_printable_runs_from_binary(tmp_path):
    # A real legacy .doc is a binary OLE file; simulate the shape (readable
    # text runs interleaved with binary control bytes) without needing a
    # real Word installation to generate one.
    blob = b"\x00\x01\xfe\xffPython Data Engineer\x00\x00\x02\xffDatabricks SQL Terraform\xff\xfe"
    path = tmp_path / "resume.doc"
    path.write_bytes(blob)
    text = extract_doc_text(path)
    assert "Python Data Engineer" in text
    assert "Databricks SQL Terraform" in text


def test_extract_doc_text_on_empty_file_returns_empty_string(tmp_path):
    path = tmp_path / "empty.doc"
    path.write_bytes(b"\x00\x00\x00")
    assert extract_doc_text(path) == ""


def test_extract_resume_text_dispatches_by_extension(tmp_path):
    docx_path = _make_docx(tmp_path, ["Python SQL"])
    assert "Python SQL" in extract_resume_text(docx_path, "resume.docx")

    doc_path = tmp_path / "legacy.doc"
    doc_path.write_bytes(b"\x00Java Spring Boot\x00")
    assert "Java Spring Boot" in extract_resume_text(doc_path, "legacy.doc")


def test_extract_resume_text_rejects_unsupported_extension(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("irrelevant")
    with pytest.raises(ValueError):
        extract_resume_text(path, "notes.txt")
